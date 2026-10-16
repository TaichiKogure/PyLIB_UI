"""Rebuild offline reference curves from the acquired public sources.

Development dependency: openpyxl. No macros/OLE code is executed. The EMF
reader extracts only pen selections and POLYLINE16 records from the workbook.
Sources are cached in data/reference/cathodes/sources (not runtime inputs).
"""
import csv
import hashlib
import json
import math
import struct
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT / "data/reference/cathodes/sources"
DESTINATION = ROOT / "pylibstudy/reference/cathodes.v1.json"
ACQUISITION = ROOT / "pylibstudy/reference/acquisition.json"


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def source(file, url, **extra):
    path = SOURCES / file
    return dict(file=file, url=url, sha256=digest(path), bytes=path.stat().st_size,
                retrieved_date="2026-10-16", **extra)


def simplify(points, tolerance=0.0005):
    """Retain endpoints; vertical interpolation error <= tolerance V at every knot."""
    keep = {0, len(points) - 1}
    stack = [(0, len(points) - 1)]
    while stack:
        a, b = stack.pop()
        qa, va = points[a]
        qb, vb = points[b]
        errors = [(abs(v - (va + (vb - va) * (q - qa) / (qb - qa))), i)
                  for i, (q, v) in enumerate(points[a + 1:b], a + 1)]
        if errors:
            error, i = max(errors)
            if error > tolerance:
                keep.add(i)
                stack.extend(((a, i), (i, b)))
    return [points[i] for i in sorted(keep)]


def branch(points, method, *, capacity=True, extra=None):
    # Plotted EMF coordinates have quantized x and can retrace a vertical tail.
    # First occurrence of each advancing x preserves the first cutoff crossing.
    advancing = []
    for q, v in points:
        if not math.isfinite(q + v):
            raise ValueError("Nonfinite source value")
        if not advancing or q > advancing[-1][0]:
            advancing.append((q, v))
    if len(advancing) < 3:
        raise ValueError("Insufficient source points")
    compact = simplify(advancing)
    q0, q1 = compact[0][0], compact[-1][0]
    return {"points": [[round((q - q0) / (q1 - q0), 12), round(v, 8)] for q, v in compact],
            "specific_capacity_mah_g": round(q1 - q0, 8) if capacity else None,
            "voltage_range_v": [min(v for q, v in advancing), max(v for q, v in advancing)],
            "processing": {"method": method, "source_point_count": len(points),
                           "advancing_point_count": len(advancing), "point_count": len(compact),
                           "duplicate_or_retraced_x_removed": len(points) - len(advancing),
                           "duplicate_policy": "first value at each strictly advancing capacity coordinate",
                           "linear_simplification_max_error_v": 0.0005, **(extra or {})}}


def emf_curves(data):
    objects, stack, pen, chains = {}, [], None, []
    offset = 0
    while offset < len(data):
        kind, size = struct.unpack_from("<II", data, offset)
        if size < 8 or offset + size > len(data):
            raise ValueError("Malformed EMF record")
        r = data[offset:offset + size]
        if kind == 38:
            handle, style, x, y, color = struct.unpack_from("<IIiiI", r, 8)
            objects[handle] = ("pen", color)
        elif kind == 95:
            objects[struct.unpack_from("<I", r, 8)[0]] = ("pen", struct.unpack_from("<I", r, 44)[0])
        elif kind in (39, 82):
            objects[struct.unpack_from("<I", r, 8)[0]] = ("other", None)
        elif kind == 37:
            obj = objects.get(struct.unpack_from("<I", r, 8)[0])
            if obj and obj[0] == "pen":
                pen = obj[1]
        elif kind == 33:
            stack.append(pen)
        elif kind == 34 and stack:
            pen = stack.pop()
        elif kind == 87:
            count = struct.unpack_from("<I", r, 24)[0]
            pts = [struct.unpack_from("<hh", r, 28 + j * 4) for j in range(count)]
            if count > 2:
                if chains and chains[-1][0] == pen and chains[-1][1][-1] == pts[0]:
                    chains[-1][1].extend(pts[1:])
                else:
                    chains.append((pen, pts))
        offset += size
    return [(tuple((color >> k) & 255 for k in (0, 8, 16)), pts)
            for color, pts in chains if len(pts) > 1000]


def main():
    curves = []
    acquisition = json.loads(ACQUISITION.read_text(encoding="utf-8"))
    pybep = acquisition["pybep"]
    for chemistry, filename, doi, composition in [
        ("NMC111", "NMC111-10.1149_2.062204jes.txt", "10.1149/2.062204jes", "LiNi1/3Mn1/3Co1/3O2"),
        ("NMC523", "NMC523-10.1149_2.1701713jes.txt", "10.1149/2.1701713jes", "LiNi0.5Mn0.3Co0.2O2")]:
        rows = [tuple(map(float, line.split())) for line in (SOURCES / filename).read_text().splitlines()]
        a = branch(rows, "author_repository_processed_ocp", capacity=False,
                   extra={
                       "coordinate": "relative reference coordinate; increasing voltage defines charged end; not absolute Li stoichiometry"})
        curves.append(dict(id=chemistry.lower() + "_pybep", chemistry=chemistry, composition=composition,
                           label=chemistry + " / PyBEP加工OCP", kind="processed_ocp", temperature_c=None,
                           specific_current_ma_g=None, cycle=None, license="CC BY-NC-SA 4.0 (PyBEP data)",
                           citation=dict(
                               title="PyBEP: An Open-Source Tool for Electrode Potential Determination from Battery OCV Measurements",
                               url="https://github.com/JonPisek/PyBEP", doi="10.3390/batteries11080295",
                               original_doi=doi),
                           sources=[source(filename, next(v["url"] for v in pybep if v.get("file") == filename))],
                           notes=["著者リポジトリの加工OCP。原計測データではありません。",
                                  "充電方向の基準OCPを両方向に使用。独立した充放電ヒステリシスはありません。",
                                  "横軸は収録範囲内の相対状態に変換。温度・比容量はこのファイルだけでは同定できません。"],
                           branches={"charge": a}, requested_max_v=4.4))

    rows = [(1 - float(r["sto"]), float(r["p"])) for r in
            csv.DictReader((SOURCES / "NMC622_OCP_Delithiation.csv").open())]
    a = branch(sorted(rows), "published_ocp_csv", capacity=False,
               extra={
                   "coordinate": "1 - source sto; normalized electrode reference coordinate, not absolute stoichiometry"})
    members = acquisition["wan2026_remote_zip"]
    wan_url = "https://zenodo.org/records/20259651"
    curves.append(dict(id="nmc622_wan2026", chemistry="NMC622", composition="LiNi0.6Mn0.2Co0.2O2",
                       label="NMC622 / Wan 2026 OCP", kind="published_ocp", temperature_c=None,
                       specific_current_ma_g=None, cycle=None, license="CC BY 4.0",
                       citation=dict(
                           title="Managing silicon burn-out via onboard material diagnostics for durable high-energy density batteries",
                           doi="10.1016/j.joule.2026.102531", url=wan_url),
                       sources=[source("NMC622_OCP_Delithiation.csv", wan_url,
                                       archive_member="data/cell_ocp/NMC622_OCP_Delithiation.csv",
                                       archive_extraction=members)],
                       notes=["著者公開の脱リチウムOCP CSV。充放電とも同じOCPを使用します。",
                              "CSVのstoは正規化基準座標として扱います。絶対Li量は推定しません。",
                              "上端は約4.300 V。4.4 Vまでの曲線は未収録です。"],
                       branches={"charge": a}, requested_max_v=4.4))

    wb = openpyxl.load_workbook(SOURCES / "nmc8112023-source.xlsx", data_only=True, read_only=True)
    rows = list(wb["Figure 4b"].iter_rows(values_only=True))
    branches = {}
    for mode, col in (("charge", 0), ("discharge", 2)):
        points = [(r[col], r[col + 1]) for r in rows if
                  isinstance(r[col], (int, float)) and isinstance(r[col + 1], (int, float))]
        branches[mode] = branch(points, "publisher_source_data_xlsx",
                                extra={"sheet": "Figure 4b", "columns_1_based": [col + 1, col + 2],
                                       "rows_1_based": [1, len(points)]})
    curves.append(dict(id="nmc811_2023_cycle4", chemistry="NMC811", composition="LiNi0.8Mn0.1Co0.1O2",
                       label="NMC811 / 25°C・4th / 4.6 V系", kind="finite_current", temperature_c=25,
                       specific_current_ma_g=100, cycle=4, license="CC BY 4.0",
                       citation=dict(
                           title="Anion-enrichment interface enables high-voltage anode-free lithium metal batteries",
                           doi="10.1038/s41467-023-36853-x", url="https://www.nature.com/articles/s41467-023-36853-x"),
                       sources=[source("nmc8112023-source.xlsx",
                                       "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-023-36853-x/MediaObjects/41467_2023_36853_MOESM2_ESM.xlsx")],
                       notes=["Fig.4bの最初の組は4th（図中凡例で照合）。0.5C、1C=200 mA/g。",
                              "Li金属半電池。1 M LiBF4 + 1 M LiDFOB tFEP/FECという特定の電解液での測定です。",
                              "3.0–4.6 V試験。有限電流の分極を含み、平衡OCPではありません。"],
                       branches=branches, requested_max_v=4.4))
    wb.close()

    wb = openpyxl.load_workbook(SOURCES / "lco2018-figure4.xlsx", data_only=True, read_only=True)
    rows = list(wb["Figure 4c"].iter_rows(min_row=3, values_only=True))
    for col, variant, label in [(3, "bare", "Bare LCO"), (27, "laf2", "2% LAF-LCO")]:
        values = [(r[col], r[col + 1]) for r in rows if
                  isinstance(r[col], (int, float)) and isinstance(r[col + 1], (int, float))]
        split = next(i for i in range(1, len(values)) if values[i][0] < values[i - 1][0] - 1)
        branches = {mode: branch(v, "publisher_source_data_xlsx",
                                 extra={"sheet": "Figure 4c", "columns_1_based": [col + 1, col + 2],
                                        "capacity_reset_index_0_based": split})
                    for mode, v in [("charge", values[:split]), ("discharge", values[split:])]}
        curves.append(dict(id="lco2018_" + variant + "_cycle3", chemistry="LCO", composition="LiCoO2",
                           label=label + " / 3rd / 4.6 V系",
                           kind="finite_current", temperature_c=None, specific_current_ma_g=27.4, cycle=3,
                           license="CC BY 4.0",
                           citation=dict(
                               title="Electrochemical surface passivation of LiCoO2 particles at ultrahigh voltage and its applications in lithium-based batteries",
                               doi="10.1038/s41467-018-07296-6",
                               url="https://pmc.ncbi.nlm.nih.gov/articles/PMC6249257/"),
                           sources=[source("lco2018-figure4.xlsx",
                                           "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-018-07296-6/MediaObjects/41467_2018_7296_MOESM4_ESM.xlsx")],
                           notes=["Fig.4cの3rd。室温（数値温度は未指定）、Li金属半電池、3.0–4.6 V。",
                                  "表の容量単位表記の誤記 mAg/g は本文・図に従い mAh/g と解釈。",
                                  label + "を独立管理。表面処理の違いを汎用LCOの特性に置換しません。"],
                           branches=branches, requested_max_v=4.55))
    wb.close()

    with zipfile.ZipFile(SOURCES / "lmfp2022-source.xlsx") as z:
        for chemistry, member, x0, x160, color, figure in [
            ("LMFP", "xl/media/image29.emf", 1117, 3541, (84, 153, 178), "Fig.6b"),
            ("LFP", "xl/media/image35.emf", 1101, 3270, (156, 132, 190), "Fig.6h")]:
            emf = z.read(member)
            paths = [pts for rgb, pts in emf_curves(emf) if rgb == color]
            assert len(paths) == 2
            branches = {}
            for mode, pts in zip(("charge", "discharge"), paths):
                # Axis ticks: y=210 => 4.5 V, y=2364 => 1.5 V.
                data = [((x - x0) * 160 / (x160 - x0), 4.5 - (y - 210) * 3 / 2154) for x, y in pts]
                branches[mode] = branch(data, "vector_figure_digitized", extra={"figure": figure, "emf_member": member,
                                                                                "emf_sha256": hashlib.sha256(
                                                                                    emf).hexdigest(),
                                                                                "pen_rgb": list(color), "cycle": 3,
                                                                                "axis_calibration": {"q0_x": x0,
                                                                                                     "q160_x": x160,
                                                                                                     "v4_5_y": 210,
                                                                                                     "v1_5_y": 2364},
                                                                                "coordinate_resolution": {
                                                                                    "capacity_mah_g": 160 / (x160 - x0),
                                                                                    "voltage_v": 3 / 2154},
                                                                                "instrument_raw": False})
            curves.append(dict(id=chemistry.lower() + "_2022_fmc_cycle3", chemistry=chemistry,
                               composition="LiMn0.6Fe0.4PO4" if chemistry == "LMFP" else "LiFePO4",
                               label=chemistry + " / 70°C・3rd / FMC-ASPE", kind="vector_digitized", temperature_c=70,
                               specific_current_ma_g=70, cycle=3, license="CC BY 4.0",
                               citation=dict(
                                   title="Rational design of a topological polymeric solid electrolyte for high-performance all-solid-state alkali metal batteries",
                                   doi="10.1038/s41467-022-31792-5",
                                   url="https://www.nature.com/articles/s41467-022-31792-5"),
                               sources=[source("lmfp2022-source.xlsx",
                                               "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-022-31792-5/MediaObjects/41467_2022_31792_MOESM2_ESM.xlsx",
                                               member=member)],
                               notes=[
                                   figure + "の3rd曲線。公開Excelに埋め込まれたEMFの線座標を抽出。計測器の生数値ではありません。",
                                   "70°C・70 mA/g、FMC-ASPE-Li固体電解質／Li金属半電池。室温・液系への換算はしません。",
                                   "同じ横座標の最初の値を採用。軸の量子化と図の線幅に由来する誤差があります。"],
                               branches=branches, requested_max_v=4.2 if chemistry == "LMFP" else 3.7))

    # Graphite is a separate OCP input, not a measured full-cell pairing.
    values = [(float(r["sto"]), float(r["p"])) for r in
              csv.DictReader((SOURCES / "Graphite_OCP_Lithiation_raw.csv").open())]
    grouped = {}
    for x, v in values:
        grouped.setdefault(x, []).append(v)
    graphite = dict(id="graphite_wan2026", label="Graphite / Wan 2026 lithiation OCP", license="CC BY 4.0",
                    citation={"url": wan_url, "doi": "10.1016/j.joule.2026.102531"},
                    source=source("Graphite_OCP_Lithiation_raw.csv", wan_url,
                                  archive_member="data/cell_ocp/Graphite_OCP_Lithiation_raw.csv"),
                    points=[[x, sum(v) / len(v)] for x, v in sorted(grouped.items())],
                    processing={"method": "equal-coordinate mean; sorted by source sto", "source_count": len(values),
                                "point_count": len(grouped)},
                    notes=["黒鉛リチウム化OCPの公開基準座標。絶対化学量論組成とは同一視しません。",
                           "両方向に同じOCPを使用。負極の分極・ヒステリシスは未計算です。"])
    authors = {
        "10.3390/batteries11080295": [
            "University of Ljubljana, Faculty of Mechanical Engineering, LICeM (PyBEP rights holder and contributors)"],
        "10.1016/j.joule.2026.102531": ["Zhiwen Wan", "Andrew Weng", "Sravan Pannala", "Hamidreza Movahedi",
                                        "Taylor R. Garrick", "Gregory J. Offer", "Jason B. Siegel",
                                        "Anna G. Stefanopoulou"],
        "10.1038/s41467-018-07296-6": ["Jiawei Qian", "Lei Liu", "Jixiang Yang", "Siyuan Li", "Xiao Wang",
                                       "Houlong L. Zhuang", "Yingying Lu"],
        "10.1038/s41467-022-31792-5": ["Yun Su", "Xiaohui Rong", "Ang Gao", "Yuan Liu", "Jianwei Li", "Minglei Mao",
                                       "Xingguo Qi", "Guoliang Chai", "Qinghua Zhang", "Liumin Suo", "Lin Gu",
                                       "Hong Li", "Xuejie Huang", "Liquan Chen", "Binyuan Liu", "Yong-Sheng Hu"],
        "10.1038/s41467-023-36853-x": ["Minglei Mao", "Xiao Ji", "Qiyu Wang", "Zejing Lin", "Meiying Li", "Tao Liu",
                                       "Chengliang Wang", "Yong-Sheng Hu", "Hong Li", "Xuejie Huang", "Liquan Chen",
                                       "Liumin Suo"],
    }
    for c in curves:
        c["citation"]["authors"] = authors[c["citation"]["doi"]]
        c["license_url"] = "https://creativecommons.org/licenses/by-nc-sa/4.0/" if c[
                                                                                       "kind"] == "processed_ocp" else "https://creativecommons.org/licenses/by/4.0/"
        c["modifications"] = "容量座標の正規化、枝の分離、同一・後退座標の処理、0.5mV以下の線形節点削減。手順は各枝のprocessing参照。"
    graphite["citation"]["authors"] = authors["10.1016/j.joule.2026.102531"]
    graphite["license_url"] = "https://creativecommons.org/licenses/by/4.0/"
    document = dict(schema_version="1.0", catalog_version="2026-10-16.1", voltage_reference="Li/Li+",
                    curves=curves, negative_reference=graphite,
                    exclusions=[{"source": "RepOD 10.18150/JKHCAM",
                                 "reason": "取得したNMC111/622/811ファイルに時間・容量・サンプリング間隔がなく、容量曲線へ変換不可"},
                                {"source": "Zenodo 15012305",
                                 "reason": "アーカイブ索引は主にMPET計算とLFP実験。LMFPの実測容量曲線として採用せず"}])
    DESTINATION.parent.mkdir(exist_ok=True)
    DESTINATION.write_text(json.dumps(document, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    for c in curves:
        print(c['id'], {k: (len(b['points']), b['specific_capacity_mah_g'], b['voltage_range_v']) for k, b in
                        c['branches'].items()})


if __name__ == "__main__":
    main()
