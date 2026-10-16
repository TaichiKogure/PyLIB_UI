"""Check acquired source hashes and voltage error at every retained raw coordinate.

Run with openpyxl available. This checks extraction/numerics, not experimental
accuracy. The original workbook cache is optional for the app, required here.
"""
import csv
import hashlib
import json
import sys
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from pylibstudy.cathodes import catalog, interpolate
from scripts.build_cathode_catalog import emf_curves


def main():
    data = catalog()
    cache = ROOT / 'data/reference/cathodes/sources'
    sources = {s['file']: s for c in data['curves'] for s in c['sources']}
    negative = data['negative_reference']['source']
    sources[negative['file']] = negative
    report = dict(catalog_version=data['catalog_version'], catalog_sha256=data['sha256'], source_hashes=[], branches=[],
                  boundary='Extraction and numerical checks only; no target-cell validation or instrument-accuracy claim.')
    for file, metadata in sources.items():
        raw = (cache / file).read_bytes()
        actual = hashlib.sha256(raw).hexdigest()
        assert actual == metadata['sha256'], file
        assert len(raw) == metadata['bytes'], file
        report['source_hashes'].append(dict(file=file, sha256=actual, bytes=len(raw), verified=True))

    def check(curve, mode, raw):
        # Independently reconstruct the advancing-capacity subsequence and
        # evaluate the app's reference interpolation at every original knot.
        advancing = []
        for q, v in raw:
            if not advancing or q > advancing[-1][0]:
                advancing.append((q, v))
        q0, q1 = advancing[0][0], advancing[-1][0]
        b = curve['branches'][mode]
        error = max(abs(v - interpolate(b['points'], (q - q0) / (q1 - q0))) for q, v in advancing)
        assert error <= .00050002, (curve['id'], mode, error)
        if b['specific_capacity_mah_g'] is not None:
            assert abs(q1 - q0 - b['specific_capacity_mah_g']) < 1e-7
        report['branches'].append(dict(curve_id=curve['id'], mode=mode, original_points=len(raw),
                                       advancing_points=len(advancing), retained_points=len(b['points']),
                                       max_interpolation_error_v=error, capacity_mah_g=b['specific_capacity_mah_g']))

    by_id = {c['id']: c for c in data['curves']}
    for chemistry in ['nmc111', 'nmc523']:
        c = by_id[chemistry + '_pybep']
        check(c, 'charge',
              [tuple(map(float, row.split())) for row in (cache / c['sources'][0]['file']).read_text().splitlines()])
    with (cache / 'NMC622_OCP_Delithiation.csv').open() as f:
        check(by_id['nmc622_wan2026'], 'charge',
              sorted((1 - float(r['sto']), float(r['p'])) for r in csv.DictReader(f)))

    wb = openpyxl.load_workbook(cache / 'nmc8112023-source.xlsx', read_only=True, data_only=True)
    rows = list(wb['Figure 4b'].iter_rows(values_only=True))
    for mode, col in [('charge', 0), ('discharge', 2)]:
        check(by_id['nmc811_2023_cycle4'], mode, [(r[col], r[col + 1]) for r in rows if
                                                  isinstance(r[col], (int, float)) and isinstance(r[col + 1],
                                                                                                  (int, float))])
    wb.close()
    wb = openpyxl.load_workbook(cache / 'lco2018-figure4.xlsx', read_only=True, data_only=True)
    rows = list(wb['Figure 4c'].iter_rows(min_row=3, values_only=True))
    for col, variant in [(3, 'bare'), (27, 'laf2')]:
        a = [(r[col], r[col + 1]) for r in rows if
             isinstance(r[col], (int, float)) and isinstance(r[col + 1], (int, float))]
        reset = next(i for i in range(1, len(a)) if a[i][0] < a[i - 1][0] - 1)
        for mode, raw in [('charge', a[:reset]), ('discharge', a[reset:])]:
            check(by_id['lco2018_' + variant + '_cycle3'], mode, raw)
    wb.close()
    with zipfile.ZipFile(cache / 'lmfp2022-source.xlsx') as z:
        for chemistry in ['lmfp', 'lfp']:
            c = by_id[chemistry + '_2022_fmc_cycle3']
            for mode, b in c['branches'].items():
                p = b['processing']
                emf = z.read(p['emf_member'])
                assert hashlib.sha256(emf).hexdigest() == p['emf_sha256']
                paths = [pts for color, pts in emf_curves(emf) if list(color) == p['pen_rgb']]
                pts = paths[0 if mode == 'charge' else 1]
                axes = p['axis_calibration']
                check(c, mode, [((x - axes['q0_x']) * 160 / (axes['q160_x'] - axes['q0_x']),
                                 4.5 - (y - axes['v4_5_y']) * 3 / (axes['v1_5_y'] - axes['v4_5_y'])) for x, y in pts])
    report['maximum_voltage_simplification_error_v'] = max(b['max_interpolation_error_v'] for b in report['branches'])
    report['positive_branch_count'] = len(report['branches'])
    report['positive_point_count'] = sum(b['retained_points'] for b in report['branches'])
    report['negative_point_count'] = len(data['negative_reference']['points'])
    path = ROOT / 'docs/evidence/cathode-curves/source-quality.json'
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: v for k, v in report.items() if k not in ['branches', 'source_hashes']}, indent=2))
    print(f"Verified {len(sources)} source files; {path}")


if __name__ == '__main__':
    main()
